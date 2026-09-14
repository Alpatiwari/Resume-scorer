"""
Extracts raw text from resume files (PDF, DOCX).

This is step 3 of the pipeline (see docs): text extraction & parsing.
Structuring the raw text into sections (skills, experience, education)
happens later in nlp_extractor.py — this module only gets clean text out
of the file.
"""
from pathlib import Path

import pdfplumber
from docx import Document


class UnsupportedFileTypeError(Exception):
    pass


class EmptyResumeError(Exception):
    """Raised when a file parses but yields no usable text (e.g. a scanned image PDF)."""
    pass


def extract_text(file_path: Path) -> str:
    suffix = file_path.suffix.lower()

    if suffix == ".pdf":
        text = _extract_pdf_text(file_path)
    elif suffix == ".docx":
        text = _extract_docx_text(file_path)
    else:
        raise UnsupportedFileTypeError(f"Unsupported file type: {suffix}")

    cleaned = text.strip()
    if not cleaned:
        raise EmptyResumeError(
            f"No extractable text found in {file_path.name}. "
            "It may be a scanned image without OCR."
        )
    return cleaned


def _extract_pdf_text(file_path: Path) -> str:
    pages_text = []
    with pdfplumber.open(file_path) as pdf:
        for page in pdf.pages:
            page_text = page.extract_text() or ""
            pages_text.append(page_text)
    return "\n".join(pages_text)


def _extract_docx_text(file_path: Path) -> str:
    doc = Document(file_path)
    paragraphs = [p.text for p in doc.paragraphs]

    # Table content (skills tables, etc.) is common in resumes and easy to miss
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    paragraphs.append(cell.text)

    return "\n".join(paragraphs)
